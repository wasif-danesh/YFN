"""Build the schema-design sample offline from checksummed official snapshots.

This is a sample investigation, not the production pipeline or approved scoring.
Python dependencies are in requirements-sample.txt. Run without Python -O because
assertions deliberately stop the sample build on validation failures.
"""
import csv
from collections import Counter
from decimal import Decimal
import hashlib
import io
import json
import math
from pathlib import Path
import platform
import sqlite3
import zipfile

import openpyxl
import numpy
import pyproj
from pyproj import Transformer, Geod
import shapefile
import shapely
from shapely.geometry import shape, Polygon, MultiPolygon
from shapely.geometry.polygon import orient
from shapely.ops import transform, unary_union
from shapely.validation import make_valid

from acquire_greater_melbourne import BASE, RAW
from melbourne_spatial import SpatialSources, geodesic_area
from melbourne_source_checks import check_sources

OUT = BASE / 'curated'
AUDIT = BASE / 'audit'
AREA_CRS = 7855  # GDA2020 / MGA zone 55; appropriate for Melbourne.
CATEGORIES = ['Parks and gardens','Natural and semi-natural open space',
              'Recreation corridor','Sportsfields and organised recreation']
QUICKSTATS = {'206041117':365,'213031348':355,'212051568':401,'212051567':380}
METHOD_VERSION = 'greater-melbourne-sample-v1-not-approved-for-publication'
TOLERANCE_M2 = 0.01  # Numerical topology tolerance, NOT a source coverage threshold.


def read_json(path):
    return json.loads(path.read_text())


def write_json(path, data):
    path.write_text(json.dumps(data,indent=2,ensure_ascii=False,allow_nan=False)+'\n')


def write_csv(path, rows):
    assert rows
    with path.open('w',newline='') as f:
        w = csv.DictWriter(f,fieldnames=list(rows[0]),lineterminator='\n')
        w.writeheader()
        w.writerows(rows)


def growth(start, end):
    return None if start is None or end is None or start <= 0 else 100*(end-start)/start


def per_resident(area, population):
    return None if area is None or population is None or population <= 0 else area/population


def esri_polygon(rings):
    """Interpret Esri clockwise shells and counterclockwise holes explicitly."""
    shells = []
    holes = []
    for ring in rings:
        p = Polygon(ring)
        if p.exterior.is_ccw:
            holes.append(p)
        else:
            shells.append(p)
    assert shells, 'Esri polygon has no clockwise exterior'
    assigned = [[] for _ in shells]
    for hole in holes:
        parents = [(s.area,i) for i,s in enumerate(shells) if s.covers(hole.representative_point())]
        assert parents, 'Unassigned interior ring'
        assigned[min(parents)[1]].append(list(hole.exterior.coords))
    polygons = [Polygon(s.exterior.coords,assigned[i]) for i,s in enumerate(shells)]
    return polygons[0] if len(polygons)==1 else MultiPolygon(polygons)


def polygonal(g):
    if g.geom_type in ['Polygon','MultiPolygon']:
        return g
    return unary_union([polygonal(p) for p in getattr(g,'geoms',[]) if p.geom_type in ['Polygon','MultiPolygon','GeometryCollection']])


def projected(g, transformer, repairs, source, feature_id):
    p = transform(transformer.transform,g)
    if not p.is_valid:
        before = p.area
        p = polygonal(make_valid(p))
        repairs.append({'source':source,'feature_id':str(feature_id),'area_before_m2':before,'area_after_m2':p.area})
    assert p.is_valid and not p.is_empty
    return p


def main():
    OUT.mkdir(exist_ok=True)
    AUDIT.mkdir(exist_ok=True)
    acquisition = read_json(BASE/'acquisition.json')
    for file, record in acquisition.items():
        assert hashlib.sha256((RAW/file).read_bytes()).hexdigest()==record['sha256'],file
    to_area = Transformer.from_crs(7844,AREA_CRS,always_xy=True,allow_ballpark=False)
    parks_to_area = Transformer.from_crs(3857,AREA_CRS,always_xy=True,allow_ballpark=False,only_best=True)
    to_lonlat = Transformer.from_crs(AREA_CRS,7844,always_xy=True)
    geod = Geod(ellps='GRS80')
    repairs = []

    # Re-extract identities and geometries from the original ABS shapefile.
    z = zipfile.ZipFile(RAW/'SA2_2021_AUST_SHP_GDA2020.zip')
    parts = {suffix:io.BytesIO(z.read(next(n for n in z.namelist() if n.endswith('.'+suffix)))) for suffix in ['shp','shx','dbf']}
    prj = z.read(next(n for n in z.namelist() if n.endswith('.prj'))).decode()
    assert pyproj.CRS.from_wkt(prj).to_epsg()==7844
    features = []
    for sr in shapefile.Reader(**parts).iterShapeRecords():
        d = sr.record.as_dict()
        if d['GCC_CODE21']=='2GMEL':
            assert sr.shape.shapeType != 0, d
            features.append({'type':'Feature','properties':d,'geometry':sr.shape.__geo_interface__})
    assert features
    features.sort(key=lambda f:f['properties']['SA2_CODE21'])
    assert len({f['properties']['SA2_CODE21'] for f in features})==len(features)
    write_json(BASE/'selected-sa2.geojson',{'type':'FeatureCollection','features':features})

    census_z = zipfile.ZipFile(RAW/'2021_GCP_SA2_for_VIC_short-header.zip')
    g02_name = next(n for n in census_z.namelist() if n.endswith('G02_VIC_SA2.csv'))
    census_rows = list(csv.DictReader(io.StringIO(census_z.read(g02_name).decode('utf-8-sig'))))
    assert len({r['SA2_CODE_2021'] for r in census_rows})==len(census_rows)
    census = {r['SA2_CODE_2021']:r for r in census_rows}
    wb = openpyxl.load_workbook(RAW/'32180DS0003_2001-25.xlsx',read_only=True,data_only=True)
    pop_rows = list(wb['Table 1'].iter_rows(max_col=35,values_only=True))
    years = pop_rows[4]
    year_cols = {y:years.index(y) for y in range(2020,2026)}
    population = {}
    for rownum,row in enumerate(pop_rows[6:],7):
        if isinstance(row[8],(int,float)):
            code = str(int(row[8]))
            assert code not in population
            population[code] = (rownum,row)

    sources = [
        {'source_id':'abs-census-2021-g02','publisher':'Australian Bureau of Statistics','dataset_name':'2021 Census GCP Victoria SA2 G02','url':'https://www.abs.gov.au/census/find-census-data/datapacks','reference_period':'2021 Census','boundary_year':2021,'observation_date':'2021-08-10','licence':'CC BY 4.0','limitation':'Reported occupied-private-dwelling rent; excludes rent-free, visitor-only and other non-classifiable households; not current advertised rent.'},
        {'source_id':'abs-erp-2024-25','publisher':'Australian Bureau of Statistics','dataset_name':'Regional population 2024-25; Table 1 SA2 ERP 2001-2025','url':'https://www.abs.gov.au/statistics/people/population/regional-population/2024-25','reference_period':'Annual 30 June estimates, 2001-2025; released 31 March 2026','boundary_year':2021,'observation_date':None,'licence':'CC BY 4.0','limitation':'2020-2021 final; 2022-2024 revised; 2025 preliminary. ERP differs from Census counts.'},
        {'source_id':'abs-sa2-2021','publisher':'Australian Bureau of Statistics','dataset_name':'ASGS Edition 3 SA2 2021 GDA2020 boundaries','url':'https://www.abs.gov.au/statistics/standards/australian-statistical-geography-standard-asgs/edition-3-july-2021-june-2026/access-and-downloads/digital-boundary-files','reference_period':'2021 boundary edition','boundary_year':2021,'observation_date':None,'licence':'CC BY 4.0','limitation':'SA2 boundaries do not necessarily match suburbs. Area calculations use EPSG:7855.'},
        {'source_id':'dtp-ptal-snapshot','publisher':'Victorian Department of Transport and Planning','dataset_name':'PTAL Melbourne metro; WFS snapshot','url':'https://discover.data.vic.gov.au/en_AU/dataset/public-transport-accessibility-level-ptal-melbourne-metro','reference_period':'Live layer retrieved 8 September 2026; linked December 2025 fact sheet describes 2 April 2025 08:00-09:00 services','boundary_year':None,'observation_date':None,'licence':'CC BY 4.0','limitation':'Feature-level service dates absent; mapping to the linked fact-sheet snapshot needs publisher confirmation. Access index does not measure destination reachability or all-day access.'},
        {'source_id':'vpa-open-space-snapshot','publisher':'Victorian Planning Authority','dataset_name':'VPA Draft Open Space Data; ArcGIS snapshot','url':'https://discover.data.vic.gov.au/dataset/open-space','reference_period':'Unmaintained inventory; observation date unverified; retrieved 8 September 2026','boundary_year':None,'observation_date':None,'licence':'CC BY 4.0','limitation':'Current completeness and ground access unverified. Excludes some public open-space categories under the sample filter. Source CRS EPSG:3857; datum transformation accuracy stated as 3 m.'},
    ]
    method_rows = [
        {'method_id':'census-g02-direct-v2','approval_status':'source_definition','definition':'Read Median_rent_weekly by SA2_CODE_2021. Preserve original tokens in the audit. The three zero entries are withheld as unavailable: zero is outside the 2021 RNTD scope and their QuickStats pages report no/very low population.'},
        {'method_id':'erp-growth-2020-2025-v1','approval_status':'sample_only','definition':'100*(ERP_2025-ERP_2020)/ERP_2020 from one release; nonpositive/missing baseline returns null.'},
        {'method_id':'ptal-area-mean-sample-v1','approval_status':'sample_only','definition':'sum(intersection_m2*sum_ai_8_9)/sum(valid_intersection_m2); EPSG:7855. No percentile, rating or coverage eligibility threshold.'},
        {'method_id':'open-space-filter-sample-v1','approval_status':'sample_only','definition':json.dumps({'OS_TYPE':'Public open space','OS_ACCESS':'Open','OS_STATUS':'Existing','OS_CATEGOR':CATEGORIES,'aggregation':'union selected polygons, intersect SA2, divide square metres by 2025 ERP','crs':'EPSG:7855'})},
    ]
    indicators = [
        {'indicator_key':'rent_weekly','label':'Median weekly rent reported in 2021 Census','unit':'AUD/week'},
        {'indicator_key':'transport_access','label':'Area-weighted morning access index (sample method)','unit':'access_index'},
        {'indicator_key':'population_growth','label':'Estimated population change 2020-2025','unit':'percent'},
        {'indicator_key':'green_space_per_resident','label':'Selected public open space per resident (sample filter)','unit':'m2/person'},
    ]
    areas=[]; history=[]; observations=[]; links=[]; components=[]; summaries=[]
    ptal_audit=[]; parks_audit=[]; spatial=[]; census_extract=[]; pop_extract=[]
    boundary_codes={f['properties']['SA2_CODE21'] for f in features}
    erp_codes={c for c,(_,r) in population.items() if r[2]=='2GMEL'}
    assert boundary_codes==erp_codes, (boundary_codes-erp_codes,erp_codes-boundary_codes)
    assert boundary_codes<=set(census)
    quickstats_checks,population_totals=check_sources(RAW,census,wb,population,year_cols,boundary_codes)
    write_json(AUDIT/'rent-cross-checks.json',quickstats_checks)
    write_json(AUDIT/'population-total-cross-checks.json',population_totals)
    source_index=SpatialSources(RAW,to_area,parks_to_area,repairs)

    for area_index,f in enumerate(features,1):
        a = f['properties']; code = a['SA2_CODE21']; name = a['SA2_NAME21']
        assert a['GCC_CODE21']=='2GMEL' and len(code)==9
        sa2 = projected(shape(f['geometry']),to_area,repairs,'abs-sa2',code)
        geo_area = geodesic_area(shape(f['geometry']),geod)
        assert abs(sa2.area/geo_area-1)<0.002  # Projection sanity check, not coverage policy.
        rownum,pop = population[code]
        assert pop[9]==name and pop[2]=='2GMEL'
        vals = {y:pop[c] for y,c in year_cols.items()}
        assert all(v is None or (isinstance(v,int) and v>=0) for v in vals.values())
        comparable=vals[2025] is not None and vals[2025]>0
        areas.append({'sa2_code':code,'name':name,'boundary_year':2021,'gccsa_code':a['GCC_CODE21'],
            'official_area_sq_km':a['AREASQKM21'],'calculated_area_m2':sa2.area,'area_crs':f'EPSG:{AREA_CRS}',
            'is_comparable':int(comparable),'eligibility_note':'Spatial Greater Melbourne SA2 with positive 2025 ERP; individual indicators may be unavailable.' if comparable else 'Retained for complete geography coverage; 2025 ERP is zero/missing, so excluded from the sample renter comparison set.',
            'source_id':'abs-sa2-2021'})
        for y,v in vals.items():
            history.append({'sa2_code':code,'year':y,'population':v,'reference_date':f'{y}-06-30',
                'revision_status':'final' if y<=2021 else 'preliminary' if y==2025 else 'revised','source_id':'abs-erp-2024-25',
                'source_locator':f"Table 1!{openpyxl.utils.get_column_letter(year_cols[y]+1)}{rownum}"})
        pop_extract.append({'sa2_code':code,'name':name,'worksheet_row':rownum,**{f'erp_{y}':v for y,v in vals.items()}})
        rent_token=census[code]['Median_rent_weekly']
        assert rent_token.isdigit(), (code,rent_token)
        rent=int(rent_token) if int(rent_token)>0 else None
        if code in QUICKSTATS:
            assert rent==QUICKSTATS[code], 'Independent QuickStats comparison failed'
        census_extract.append(census[code])

        # Check complete retrieval and unique IDs for the candidate bounding box.
        ptal_features=source_index.ptal_candidates(sa2)
        expected_ptal=len(ptal_features)
        clips=[]; valid_clips=[]; weighted_terms=[]; indices=[]
        for p,g in ptal_features:
            clipped=g.intersection(sa2)
            ar=clipped.area
            if ar<=0: continue
            value=p['properties']['sum_ai_8_9']
            clips.append(clipped)
            if value is not None:
                assert isinstance(value,(int,float)) and math.isfinite(value) and value>=0
                valid_clips.append(clipped); weighted_terms.append(ar*value); indices.append(value)
            ptal_audit.append({'sa2_code':code,'feature_id':p['id'],'source_file':p['source_file'],'sum_ai_8_9':value,'category_8_9':p['properties']['category_8_9'],
                'cell_area_m2':g.area,'intersection_m2':ar,'weighted_term':None if value is None else ar*value})
        valid_union=unary_union(valid_clips)
        covered=unary_union(clips).area
        valid_area=math.fsum(g.area for g in valid_clips)
        overlap=valid_area-valid_union.area
        assert abs(overlap)<TOLERANCE_M2, f'Overlapping PTAL cells in {code}: {overlap}'
        ptal=math.fsum(weighted_terms)/valid_area if valid_area>0 else None
        coverage=valid_union.area/sa2.area
        assert 0<=coverage<=1+1e-9
        if ptal is not None: assert min(indices)-1e-12<=ptal<=max(indices)+1e-12
        coverage=min(1.0,coverage)
        # Alternative Decimal calculation from per-cell audit components.
        if valid_area>0:
            decimal_mean=sum(Decimal(str(v)) for v in weighted_terms)/sum(Decimal(str(g.area)) for g in valid_clips)
            assert math.isclose(ptal,float(decimal_mean),rel_tol=1e-12,abs_tol=1e-12)

        park_features=source_index.park_candidates(sa2)
        expected_parks=len(park_features)
        eligible=[]; eligible_clips=[]; broader=[]; without_flagged_water=[]
        for p,g in park_features:
            attrs=p['attributes']; fid=attrs['FID']
            clipped=g.intersection(sa2)
            ar=clipped.area
            if ar<=0: continue
            public=attrs['OS_TYPE']=='Public open space' and attrs['OS_ACCESS']=='Open' and attrs['OS_STATUS']=='Existing'
            include=public and attrs['OS_CATEGOR'] in CATEGORIES
            if public: broader.append(clipped)
            if include:
                eligible.append(g);eligible_clips.append(clipped)
                if (attrs['WATER_BODY'] or '').strip()!='Yes':without_flagged_water.append(clipped)
            reasons=[]
            for field,value in [('OS_TYPE','Public open space'),('OS_ACCESS','Open'),('OS_STATUS','Existing')]:
                if attrs[field]!=value:reasons.append(f'{field}={attrs[field]}')
            if attrs['OS_CATEGOR'] not in CATEGORIES:reasons.append('category outside sample filter')
            parks_audit.append({'sa2_code':code,'FID':fid,'source_file':p['source_file'],'VPA_ID':attrs['VPA_ID'],'park_name':(attrs['PARK_NAME'] or '').strip(),
                'category':attrs['OS_CATEGOR'],'access':attrs['OS_ACCESS'],'land_type':attrs['OS_TYPE'],'status':attrs['OS_STATUS'],
                'water_body':(attrs['WATER_BODY'] or '').strip() or None,'coastal':(attrs['COASTAL'] or '').strip() or None,
                'included':int(include),'exclusion_reason':'; '.join(reasons) or None,'source_HA':attrs['HA'],
                'calculated_feature_m2':g.area,'intersection_m2':ar})
        park_union=unary_union(eligible).intersection(sa2)
        park_area=park_union.area
        # Independent order of operations: clip individual parcels then dissolve.
        alternate=unary_union(eligible_clips)
        assert park_union.symmetric_difference(alternate).area<TOLERANCE_M2
        overlap_parks=math.fsum(g.area for g in eligible_clips)-park_area
        assert overlap_parks>=-TOLERANCE_M2 and park_area<=sa2.area+TOLERANCE_M2
        broad_area=unary_union(broader).area
        assert broad_area>=park_area-TOLERANCE_M2
        # A zero inventory match cannot establish absence of parks when inventory
        # coverage is unknown. Retain area=0 in diagnostics, but withhold the metric.
        green=per_resident(park_area,vals[2025]) if park_area>0 else None
        pop_growth=growth(vals[2020],vals[2025])
        projected_union_geo=transform(to_lonlat.transform,park_union)
        geodesic_parks=geodesic_area(projected_union_geo,geod)
        if park_area>0: assert abs(park_area/geodesic_parks-1)<0.002
        diagnostics={'sa2_code':code,'name':name,'sa2_area_m2':sa2.area,'geodesic_sa2_area_m2':geo_area,
            'ptal_candidate_count':expected_ptal,'ptal_intersecting_count':len(clips),'ptal_valid_count':len(valid_clips),
            'ptal_valid_covered_area_m2':valid_union.area,'ptal_coverage_fraction':coverage,'ptal_missing_coverage_m2':max(0.0,sa2.area-valid_union.area),
            'ptal_overlap_m2':max(0.0,overlap),'ptal_min':min(indices) if indices else None,'ptal_max':max(indices) if indices else None,
            'park_candidate_count':expected_parks,'park_included_feature_count':len(eligible),'park_union_area_m2':park_area,
            'park_overlap_removed_m2':max(0.0,overlap_parks),'park_geodesic_area_m2':geodesic_parks,
            'broader_public_open_space_m2':broad_area,'broader_public_open_space_m2_per_person':per_resident(broad_area,vals[2025]),
            'selected_space_excluding_explicit_water_body_yes_m2_per_person':per_resident(unary_union(without_flagged_water).area,vals[2025]),
            'open_space_inventory_coverage_fraction':None}
        spatial.append(diagnostics)

        def obs(key,value,method,period,start,end,quality,note,cov=None):
            if value is None: quality='unavailable'
            observations.append({'sa2_code':code,'indicator_key':key,'raw_value':value,'score':None,'rating':None,
                'reference_period':period,'start_year':start,'end_year':end,'quality_status':quality,'quality_note':note,
                'coverage_fraction':cov,'method_id':method})
        obs('rent_weekly',rent,'census-g02-direct-v2','2021 Census',2021,2021,'available','Historical reported rent; not current advertised rent.' if rent is not None else 'Source G02 token is 0; median withheld. 2021 RNTD excludes zero rent and the area QuickStats reports no/very low population. This does not mean free rent.')
        obs('population_growth',pop_growth,'erp-growth-2020-2025-v1','30 June 2020 to 30 June 2025',2020,2025,'available','2025 ERP is preliminary; historical estimated change, not a forecast.' if pop_growth is not None else '2020 population is zero/missing or 2025 population is missing; percentage change is undefined.')
        obs('transport_access',ptal,'ptal-area-mean-sample-v1','Morning access index; live snapshot retrieved 2026-09-08',None,None,'limited',
            'Sample area-weighted raw index over covered land only. Weighting and minimum coverage are unapproved; no approved benchmark or percentile. Layer observation date unverified.' if ptal is not None else 'No valid PTAL polygon area intersects this SA2. Outside coverage is unknown, not zero access.',coverage)
        obs('green_space_per_resident',green,'open-space-filter-sample-v1','Undated unmaintained open-space inventory / 30 June 2025 ERP',None,None,'limited',
            ('Zero/missing 2025 population; per-resident value undefined. ' if not comparable else 'No selected open-space polygon intersects this SA2; inventory coverage is unknown, so no zero is asserted. ' if park_area==0 else '')+'Sample public-open-space filter includes sports grounds and natural open space that can include river land/water. Not confirmed vegetation or current access. Inventory completeness/date unverified. Denominator is preliminary 2025 ERP.')
        for key,pairs in {
            'rent_weekly':[('abs-census-2021-g02','raw_measure'),('abs-sa2-2021','boundary')],
            'population_growth':[('abs-erp-2024-25','endpoints'),('abs-sa2-2021','boundary')],
            'transport_access':[('dtp-ptal-snapshot','raw_measure'),('abs-sa2-2021','boundary')],
            'green_space_per_resident':[('vpa-open-space-snapshot','numerator'),('abs-erp-2024-25','denominator'),('abs-sa2-2021','boundary')],
        }.items():
            for source,role in pairs:links.append({'sa2_code':code,'indicator_key':key,'source_id':source,'role':role})
        def component(key,component_key,value,unit,period,source):
            components.append({'sa2_code':code,'indicator_key':key,'component_key':component_key,'value':value,'unit':unit,'reference_period':period,'source_id':source})
        component('population_growth','population_start',vals[2020],'persons','2020-06-30','abs-erp-2024-25')
        component('population_growth','population_end',vals[2025],'persons','2025-06-30','abs-erp-2024-25')
        component('transport_access','weighted_index_area_sum',math.fsum(weighted_terms),'access_index*m2','Retrieved 2026-09-08','dtp-ptal-snapshot')
        component('transport_access','valid_intersection_area',valid_area,'m2','Retrieved 2026-09-08','dtp-ptal-snapshot')
        component('transport_access','sa2_area',sa2.area,'m2','Boundary 2021','abs-sa2-2021')
        component('green_space_per_resident','eligible_open_space_area',park_area,'m2','Observation date unverified','vpa-open-space-snapshot')
        component('green_space_per_resident','population_denominator',vals[2025],'persons','2025-06-30','abs-erp-2024-25')
        summaries.append({'sa2_code':code,'name':name,'is_comparable':int(comparable),'rent_2021_aud_week':rent,'erp_2020':vals[2020],'erp_2025':vals[2025],
            'population_growth_2020_2025_percent':pop_growth,'transport_raw_index_sample':ptal,'transport_coverage_fraction':coverage,
            'selected_open_space_m2':park_area,'selected_open_space_m2_per_person_sample':green,'transport_score':None,
            'spatial_method_status':'sample_only_not_approved_for_publication'})
        if area_index%20==0 or area_index==len(features):
            print(f'Processed {area_index}/{len(features)} SA2s: {name}; PTAL coverage {coverage:.4f}',flush=True)

    # Retain source files as a separate table: derived observations can refer to
    # several datasets, and a dataset can have several spatial query files.
    source_files=[]
    for filename,r in acquisition.items():
        sid = 'abs-census-2021-g02' if filename.startswith(('2021_GCP','quickstats','rent-dictionary')) else 'abs-erp-2024-25' if filename.endswith('.xlsx') else 'abs-sa2-2021' if filename.startswith('SA2_') else 'dtp-ptal-snapshot' if filename.startswith('ptal') else 'vpa-open-space-snapshot'
        source_files.append({'file_id':filename,'source_id':sid,'path':r['file'],'resource_url':r['url'],'sha256':r['sha256'],'downloaded_at':r['retrieved_at'],'bytes':r['bytes']})
    tables={'areas':areas,'indicators':indicators,'methods':method_rows,'observations':observations,'population_history':history,
        'data_sources':sources,'source_files':source_files,'observation_sources':links,'observation_components':components}
    diagnostic_fields=['sa2_code','ptal_candidate_count','ptal_intersecting_count','ptal_valid_count','ptal_valid_covered_area_m2',
        'ptal_coverage_fraction','park_candidate_count','park_included_feature_count','park_union_area_m2','park_overlap_removed_m2',
        'selected_space_excluding_explicit_water_body_yes_m2_per_person','open_space_inventory_coverage_fraction']
    tables['area_diagnostics']=[{k:r[k] for k in diagnostic_fields} for r in spatial]
    # Flatten statuses alongside values, preventing CSV consumers from silently
    # treating provisional or missing measurements as fully available statistics.
    by_key={(r['sa2_code'],r['indicator_key']):r for r in observations}
    for r in summaries:
        for key in ['rent_weekly','population_growth','transport_access','green_space_per_resident']:
            o=by_key[r['sa2_code'],key]
            r[key+'_quality']=o['quality_status'];r[key+'_note']=o['quality_note']
    for name,rows in tables.items():write_csv(OUT/f'{name}.csv',rows)
    write_csv(BASE/'sample-summary.csv',summaries)
    write_json(BASE/'sample-summary.json',summaries)
    write_csv(AUDIT/'census-g02-source-rows.csv',census_extract)
    write_csv(AUDIT/'population-source-rows.csv',pop_extract)
    write_csv(AUDIT/'ptal-intersections.csv',ptal_audit)
    write_csv(AUDIT/'open-space-intersections.csv',parks_audit)
    # A separate publisher-stored hectare attribute corroborates full-feature
    # area calculations. It is rounded to 0.0001 ha (1 square metre).
    ha_differences=[abs(r['calculated_feature_m2']-float(r['source_HA'])*10000) for r in parks_audit if r['source_HA'] is not None]
    ha_outliers=[r for r in parks_audit if r['source_HA'] is not None and abs(r['calculated_feature_m2']-float(r['source_HA'])*10000)>0.51]
    # Do not assume the small-sample source HA rounding agreement holds globally.
    # Preserve exceptions for inspection; geometry, not source HA, drives results.
    write_json(AUDIT/'source-hectare-exceptions.json',ha_outliers)
    write_json(AUDIT/'spatial-validation.json',spatial)
    write_json(AUDIT/'geometry-repairs.json',repairs)
    write_json(AUDIT/'field-mapping.json',{
        'area_identity':{'resource':'SA2_2021_AUST_GDA2020.dbf','code':'SA2_CODE21','name':'SA2_NAME21','gccsa':'GCC_CODE21','area':'AREASQKM21'},
        'rent':{'resource':g02_name,'code':'SA2_CODE_2021','value':'Median_rent_weekly','dictionary_cell':'G112','table':'G02','unit':'AUD/week'},
        'population':{'resource':'32180DS0003_2001-25.xlsx','sheet':'Table 1','code_column':'I','name_column':'J','year_header_row':5,'data_start_row':7,'years':year_cols,'year_column_indices_are_zero_based':True},
        'transport':{'layer':'open-data-platform:ptal_metro','numeric':'sum_ai_8_9','category':'category_8_9','geometry':'geom','crs':'EPSG:7844'},
        'open_space':{'layer':'VPA_Draft_Open_Space_Data','id':'FID','secondary_id':'VPA_ID','filter_fields':['OS_TYPE','OS_ACCESS','OS_STATUS','OS_CATEGOR'],'geometry':'rings','crs':'EPSG:3857'}})
    write_json(AUDIT/'spatial-source-inventory.json',source_index.inventory)
    manifest={'data_mode':'real','release_id':'greater-melbourne-v1','purpose':'schema_design_sample','publication_ready':False,
        'method_version':METHOD_VERSION,'boundary_year':2021,'population_start_year':2020,'population_end_year':2025,
        'transport_benchmark_count':None,'transport_score_status':'withheld_no_approved_benchmark_or_method',
        'open_space_observation_date':None,'area_count':len(areas),'observation_count':len(observations),'population_history_count':len(history),
        'geographic_scope':'All spatial ASGS 2021 SA2s classified to GCCSA 2GMEL (Greater Melbourne)',
        'comparison_area_count':sum(r['is_comparable'] for r in areas),
        'open_space_sample_categories':CATEGORIES,'source_manifest_sha256':hashlib.sha256((BASE/'acquisition.json').read_bytes()).hexdigest(),
        'runtime':{'python':platform.python_version(),'shapely':shapely.__version__,'pyproj':pyproj.__version__,'pyshp':shapefile.__version__,'openpyxl':openpyxl.__version__},
        'pipeline_files_sha256':{name:hashlib.sha256((Path(__file__).parent/name).read_bytes()).hexdigest() for name in
            ['build_greater_melbourne.py','melbourne_spatial.py','melbourne_source_checks.py','acquire_greater_melbourne.py',
             'build_real_sample.py','acquire_real_sample.py','greater_melbourne_schema.sql','requirements-greater-melbourne.txt','requirements-sample.txt']},
        'crs':{'area':'EPSG:7855','sa2_and_ptal':'EPSG:7844','open_space':'EPSG:3857','open_space_transform':parks_to_area.get_last_used_operation().description,'open_space_transform_accuracy_m':parks_to_area.get_last_used_operation().accuracy},
        'verification_scope':'All-area checksums, source identities, population endpoints, complete spatial layer counts and IDs, clipping, overlap, geodesic area sanity, database constraints; separate QuickStats spot checks. No field survey or approved production scoring.'}
    manifest['runtime']['numpy']=numpy.__version__
    write_json(BASE/'manifest.json',manifest)
    # Load actual CSV files, exercising the agreed interchange rather than using
    # an independent in-memory path to the database.
    schema=(Path(__file__).parent/'greater_melbourne_schema.sql').read_text()
    db_temp=BASE/'sample.building.sqlite'
    if db_temp.exists():db_temp.unlink()
    db=sqlite3.connect(db_temp)
    db.executescript(schema)
    with db:
        for name in ['data_sources','source_files','methods','areas','indicators','observations','population_history','observation_sources','observation_components','area_diagnostics']:
            with (OUT/f'{name}.csv').open(newline='') as f:
                records=csv.DictReader(f);fields=records.fieldnames
                query=f'INSERT INTO {name} ({",".join(fields)}) VALUES ({",".join("?" for _ in fields)})'
                db.executemany(query,[[None if v=='' else v for v in r.values()] for r in records])
        db.execute('INSERT INTO sample_release VALUES (?,?,?,?)',('greater-melbourne-v1','real',0,json.dumps(manifest)))
    assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
    assert db.execute('PRAGMA foreign_key_check').fetchall()==[]
    assert db.execute('SELECT count(*) FROM observations').fetchone()[0]==4*len(areas)
    assert db.execute('SELECT count(*) FROM population_history').fetchone()[0]==6*len(areas)
    assert db.execute('SELECT count(*) FROM observations WHERE score IS NOT NULL OR rating IS NOT NULL').fetchone()[0]==0
    assert db.execute('SELECT count(*) FROM areas WHERE typeof(sa2_code)!="text"').fetchone()[0]==0
    for code,key,value in db.execute('SELECT sa2_code,indicator_key,raw_value FROM observations'):
        expected=next(r['raw_value'] for r in observations if r['sa2_code']==code and r['indicator_key']==key)
        assert (value is None and expected is None) or (value is not None and expected is not None and math.isclose(value,expected,rel_tol=1e-12))
    db.close()
    db_temp.replace(BASE/'sample.sqlite')
    write_json(AUDIT/'validation-report.json',{'status':'passed_with_documented_limitations','table_counts':{n:len(r) for n,r in tables.items()},
        'raw_file_checksums_verified':len(acquisition),'quickstats_pages_cross_checked':len(quickstats_checks),
        'quickstats_numeric_rent_matches':sum(c['quickstats_rent'] is not None for c in quickstats_checks),
        'quickstats_zero_rent_unavailability_confirmations':sum(c['quickstats_rent'] is None for c in quickstats_checks),
        'population_annual_gccsa_total_matches':len(population_totals),'sa2_census_erp_identity_matches':len(areas),
        'source_query_counts_and_unique_ids_match':True,'geometry_repairs':len(repairs),'unique_repaired_source_features':len({(r['source'],r['feature_id']) for r in repairs}),'spatial_order_of_operations_check':True,
        'max_full_feature_area_difference_from_source_hectares_m2':max(ha_differences),'source_hectare_area_comparisons':len(ha_differences),
        'source_hectare_exceptions':len(ha_outliers),
        'quality_counts':{key:dict(Counter(r['quality_status'] for r in observations if r['indicator_key']==key)) for key in [r['indicator_key'] for r in indicators]},
        'projection_vs_geodesic_relative_tolerance':0.002,'topology_tolerance_m2':TOLERANCE_M2,'sqlite_integrity':'ok','foreign_key_violations':0,
        'limitations':['Transport weighting and open-space filter are sample methods, not production-approved.',
            'Transport and open-space observation dates are not asserted from metadata update dates.',
            'Complete API query retrieval does not prove geographic inventory completeness or current ground truth.',
            '2025 population is preliminary; rent is 2021 Census reported rent.',
            'Transport percentile and rating are intentionally null.']})


if __name__=='__main__':
    main()
